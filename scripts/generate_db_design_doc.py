from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path("database_table_design.docx")


TABLES = [
    {
        "name": "users",
        "description": "Lưu thông tin tài khoản người dùng và phân quyền hệ thống.",
        "rows": [
            ("id", "INTEGER", "Primary key, tự tăng."),
            ("full_name", "VARCHAR(255)", "Họ tên người dùng."),
            ("email", "VARCHAR(255)", "Unique. Email đăng nhập."),
            ("password_hash", "VARCHAR(255)", "Mật khẩu đã hash; không trả về API response."),
            ("phone", "VARCHAR(30)", "Nullable. Số điện thoại người dùng."),
            ("role", "ENUM", "Vai trò: CUSTOMER, ADMIN. Default CUSTOMER."),
            ("status", "ENUM", "Trạng thái: ACTIVE, LOCKED, DELETED. Default ACTIVE."),
            ("created_at", "TIMESTAMP", "Thời điểm tạo tài khoản."),
            ("updated_at", "TIMESTAMP", "Thời điểm cập nhật tài khoản gần nhất."),
        ],
    },
    {
        "name": "user_addresses",
        "description": "Lưu danh sách địa chỉ nhận hàng của người dùng.",
        "rows": [
            ("id", "INTEGER", "Primary key, tự tăng."),
            ("receiver_name", "VARCHAR(255)", "Tên người nhận hàng."),
            ("receiver_phone", "VARCHAR(30)", "Số điện thoại người nhận hàng."),
            ("address_line", "TEXT", "Địa chỉ chi tiết."),
            ("ward", "VARCHAR(100)", "Nullable. Phường/xã."),
            ("province", "VARCHAR(100)", "Nullable. Tỉnh/thành phố."),
            ("is_default", "BOOLEAN", "Default false. Đánh dấu địa chỉ mặc định."),
            ("updated_at", "TIMESTAMP", "Thời điểm cập nhật địa chỉ gần nhất."),
            ("user_id", "INTEGER", "Foreign key -> users.id. ON DELETE CASCADE."),
        ],
    },
    {
        "name": "carts",
        "description": "Giỏ hàng hiện tại của mỗi người dùng.",
        "rows": [
            ("id", "INTEGER", "Primary key, tự tăng."),
            ("created_at", "TIMESTAMP", "Thời điểm tạo giỏ hàng."),
            ("updated_at", "TIMESTAMP", "Thời điểm cập nhật giỏ hàng gần nhất."),
            ("user_id", "INTEGER", "Foreign key -> users.id. Quan hệ 1-1, ON DELETE CASCADE."),
        ],
    },
    {
        "name": "cart_items",
        "description": "Các sản phẩm trong giỏ hàng, theo đơn vị tính cụ thể.",
        "rows": [
            ("id", "INTEGER", "Primary key, tự tăng."),
            ("quantity", "INTEGER", "Số lượng sản phẩm trong giỏ."),
            ("unit_price_snapshot", "DECIMAL(12,2)", "Nullable. Giá tại thời điểm thêm vào giỏ."),
            ("cart_id", "INTEGER", "Foreign key -> carts.id. ON DELETE CASCADE."),
            ("product_id", "INTEGER", "Foreign key -> products.id. ON DELETE CASCADE."),
            ("measure_unit_id", "INTEGER", "Nullable. Foreign key -> measure_units.id. ON DELETE SET NULL."),
        ],
        "extra_note": "Unique composite: cart_id + product_id + measure_unit_id.",
    },
    {
        "name": "products",
        "description": "Thông tin thuốc/sản phẩm và nội dung y tế đã chuẩn hóa.",
        "rows": [
            ("id", "INTEGER", "Primary key, tự tăng."),
            ("name", "VARCHAR(500)", "Tên sản phẩm."),
            ("slug", "VARCHAR(700)", "Unique. Slug dùng cho URL sản phẩm."),
            ("product_type", "ENUM", "Loại sản phẩm: DRUG, SUPPLEMENT, OTHER. Default OTHER."),
            ("description", "TEXT", "Nullable. Mô tả sản phẩm."),
            ("image_url", "VARCHAR(1000)", "Nullable. URL hình ảnh sản phẩm."),
            ("is_active", "BOOLEAN", "Default true. Trạng thái hiển thị/bán."),
            ("created_at", "TIMESTAMP", "Thời điểm tạo sản phẩm."),
            ("updated_at", "TIMESTAMP", "Thời điểm cập nhật sản phẩm gần nhất."),
            ("deleted_at", "TIMESTAMP", "Nullable. Soft delete."),
            ("usage", "TEXT", "Nullable. Công dụng/cách dùng tổng quan."),
            ("dosage", "TEXT", "Nullable. Liều dùng."),
            ("adverse_effect", "TEXT", "Nullable. Tác dụng không mong muốn."),
            ("careful", "TEXT", "Nullable. Lưu ý/thận trọng khi dùng."),
            ("preservation", "TEXT", "Nullable. Cách bảo quản."),
        ],
        "extra_note": "Index: is_active + created_at; unique/index trên slug.",
    },
    {
        "name": "measure_units",
        "description": "Danh mục đơn vị tính dùng cho giá, giỏ hàng và đơn hàng.",
        "rows": [
            ("id", "INTEGER", "Primary key, tự tăng."),
            ("name", "VARCHAR(100)", "Unique. Tên đơn vị tính, ví dụ: Viên, Hộp, Gói."),
        ],
    },
    {
        "name": "product_prices",
        "description": "Bảng giá sản phẩm theo từng đơn vị tính.",
        "rows": [
            ("id", "INTEGER", "Primary key, tự tăng."),
            ("price", "DECIMAL(12,2)", "Giá bán theo đơn vị tính."),
            ("is_sell_default", "BOOLEAN", "Default false. Đơn vị bán mặc định của sản phẩm."),
            ("updated_at", "TIMESTAMP", "Thời điểm cập nhật giá gần nhất."),
            ("product_id", "INTEGER", "Foreign key -> products.id. ON DELETE CASCADE."),
            ("measure_unit_id", "INTEGER", "Foreign key -> measure_units.id. ON DELETE RESTRICT."),
        ],
        "extra_note": "Unique composite: product_id + measure_unit_id.",
    },
    {
        "name": "categories",
        "description": "Danh mục sản phẩm dạng cây cha-con.",
        "rows": [
            ("id", "INTEGER", "Primary key, tự tăng."),
            ("name", "VARCHAR(255)", "Tên danh mục."),
            ("slug", "VARCHAR(320)", "Unique. Slug danh mục."),
            ("level", "SMALLINT", "Default 1. Cấp danh mục trong cây."),
            ("is_active", "BOOLEAN", "Default true. Trạng thái hoạt động."),
            ("updated_at", "TIMESTAMP", "Thời điểm cập nhật danh mục gần nhất."),
            ("parent_id", "INTEGER", "Nullable. Foreign key -> categories.id. ON DELETE SET NULL."),
        ],
        "extra_note": "Index: slug, parent_id.",
    },
    {
        "name": "product_categories",
        "description": "Bảng trung gian gán sản phẩm vào danh mục.",
        "rows": [
            ("product_id", "INTEGER", "Composite primary key. Foreign key -> products.id. ON DELETE CASCADE."),
            ("category_id", "INTEGER", "Composite primary key. Foreign key -> categories.id. ON DELETE CASCADE."),
            ("is_primary", "BOOLEAN", "Default false. Đánh dấu danh mục chính của sản phẩm."),
        ],
        "extra_note": "Index: product_id, category_id.",
    },
    {
        "name": "orders",
        "description": "Đơn hàng được tạo sau khi checkout giỏ hàng.",
        "rows": [
            ("id", "INTEGER", "Primary key, tự tăng."),
            ("order_no", "VARCHAR(30)", "Unique. Mã đơn hàng hiển thị."),
            ("total_amount", "DECIMAL(12,2)", "Tổng tiền đơn hàng."),
            ("status", "ENUM", "Trạng thái: PENDING, PAID, CANCELLED."),
            ("note", "TEXT", "Nullable. Ghi chú đơn hàng."),
            ("created_at", "TIMESTAMP", "Thời điểm tạo đơn hàng."),
            ("updated_at", "TIMESTAMP", "Thời điểm cập nhật đơn hàng gần nhất."),
            ("user_id", "INTEGER", "Nullable. Foreign key -> users.id. ON DELETE SET NULL."),
        ],
        "extra_note": "Index: user_id.",
    },
    {
        "name": "order_items",
        "description": "Chi tiết sản phẩm trong đơn hàng, lưu snapshot tại thời điểm checkout.",
        "rows": [
            ("id", "INTEGER", "Primary key, tự tăng."),
            ("quantity", "INTEGER", "Số lượng mua."),
            ("product_name_snapshot", "VARCHAR(500)", "Tên sản phẩm tại thời điểm đặt hàng."),
            ("measure_unit_name_snapshot", "VARCHAR(100)", "Nullable. Tên đơn vị tính tại thời điểm đặt hàng."),
            ("unit_price", "DECIMAL(12,2)", "Đơn giá snapshot."),
            ("line_total", "DECIMAL(12,2)", "Thành tiền = quantity * unit_price."),
            ("order_id", "INTEGER", "Foreign key -> orders.id. ON DELETE CASCADE."),
            ("product_id", "INTEGER", "Nullable. Foreign key -> products.id. ON DELETE SET NULL."),
            ("cart_item_id", "INTEGER", "Nullable. Lưu id cart item gốc để truy vết checkout."),
        ],
        "extra_note": "Index: order_id.",
    },
    {
        "name": "payment_methods",
        "description": "Danh mục phương thức thanh toán.",
        "rows": [
            ("id", "INTEGER", "Primary key, tự tăng."),
            ("code", "VARCHAR(50)", "Unique. Mã phương thức, ví dụ: COD, VNPAY."),
            ("name", "VARCHAR(100)", "Tên hiển thị của phương thức thanh toán."),
            ("is_active", "BOOLEAN", "Default true. Trạng thái sử dụng."),
        ],
    },
    {
        "name": "payment_transactions",
        "description": "Giao dịch thanh toán phát sinh cho đơn hàng.",
        "rows": [
            ("id", "INTEGER", "Primary key, tự tăng."),
            ("amount", "DECIMAL(12,2)", "Số tiền thanh toán."),
            ("status", "ENUM", "Trạng thái: PENDING, SUCCESS, FAILED."),
            ("provider_txn_id", "VARCHAR(120)", "Nullable. Mã giao dịch/tham chiếu từ cổng thanh toán."),
            ("paid_at", "TIMESTAMP", "Nullable. Thời điểm thanh toán thành công."),
            ("created_at", "TIMESTAMP", "Thời điểm tạo giao dịch."),
            ("updated_at", "TIMESTAMP", "Thời điểm cập nhật giao dịch gần nhất."),
            ("order_id", "INTEGER", "Foreign key -> orders.id. ON DELETE CASCADE."),
            ("payment_method_id", "INTEGER", "Foreign key -> payment_methods.id. ON DELETE RESTRICT."),
        ],
        "extra_note": "Partial unique index: mỗi order chỉ có một payment SUCCESS.",
    },
    {
        "name": "product_search_embeddings",
        "description": "Vector embedding của sản phẩm dùng cho semantic search.",
        "rows": [
            ("product_id", "INTEGER", "Primary key. Foreign key -> products.id. ON DELETE CASCADE."),
            ("search_text", "TEXT", "Nội dung dùng để tạo embedding."),
            ("search_text_hash", "VARCHAR(64)", "Hash của search_text để phát hiện thay đổi."),
            ("embedding", "VECTOR(768)", "Vector embedding của sản phẩm."),
            ("embedding_model", "VARCHAR(100)", "Tên model embedding."),
            ("embedding_dimensions", "INTEGER", "Số chiều vector embedding."),
            ("created_at", "TIMESTAMP", "Thời điểm tạo embedding."),
            ("updated_at", "TIMESTAMP", "Thời điểm cập nhật embedding gần nhất."),
        ],
        "extra_note": "Index HNSW trên embedding với vector_cosine_ops.",
    },
    {
        "name": "search_query_embedding_cache",
        "description": "Cache embedding của câu truy vấn để giảm số lần gọi embedding provider.",
        "rows": [
            ("query_hash", "VARCHAR(64)", "Primary key. Hash của query_text."),
            ("query_text", "TEXT", "Nội dung truy vấn gốc."),
            ("embedding", "VECTOR(768)", "Vector embedding của truy vấn."),
            ("embedding_model", "VARCHAR(100)", "Tên model embedding."),
            ("embedding_dimensions", "INTEGER", "Số chiều vector embedding."),
            ("hit_count", "INTEGER", "Default 0. Số lần cache được dùng."),
            ("created_at", "TIMESTAMP", "Thời điểm tạo cache."),
            ("last_used_at", "TIMESTAMP", "Thời điểm cache được dùng gần nhất."),
        ],
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table_design(doc, table_def):
    heading = doc.add_paragraph()
    heading.style = doc.styles["Heading 2"]
    heading.paragraph_format.keep_with_next = True
    heading.add_run(table_def["name"])

    desc = doc.add_paragraph(table_def["description"])
    desc.style = doc.styles["Normal"]
    desc.paragraph_format.keep_with_next = True

    word_table = doc.add_table(rows=1, cols=3)
    word_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    word_table.style = "Table Grid"
    word_table.autofit = False
    set_table_width(word_table, 9360)

    widths = [Cm(4.1), Cm(4.0), Cm(8.4)]
    headers = ["Field", "Type", "Note"]
    for idx, header in enumerate(headers):
        cell = word_table.rows[0].cells[idx]
        cell.width = widths[idx]
        set_cell_shading(cell, "E8EEF5")
        set_cell_text(cell, header, bold=True, color="0B2545")
    for field, col_type, note in table_def["rows"]:
        cells = word_table.add_row().cells
        values = [field, col_type, note]
        for idx, value in enumerate(values):
            cells[idx].width = widths[idx]
            set_cell_text(cells[idx], value, bold=(idx == 0))

    if table_def.get("extra_note"):
        note_p = doc.add_paragraph()
        note_p.paragraph_format.space_before = Pt(2)
        note_p.paragraph_format.space_after = Pt(6)
        run = note_p.add_run("Ghi chú: ")
        run.bold = True
        note_p.add_run(table_def["extra_note"])
    else:
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(3)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Normal"].paragraph_format.line_spacing = 1.08

    styles["Heading 1"].font.name = "Calibri"
    styles["Heading 1"].font.size = Pt(17)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor(46, 116, 181)
    styles["Heading 1"].paragraph_format.space_before = Pt(0)
    styles["Heading 1"].paragraph_format.space_after = Pt(8)

    styles["Heading 2"].font.name = "Calibri"
    styles["Heading 2"].font.size = Pt(12.5)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].font.color.rgb = RGBColor(31, 77, 120)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(3)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("Thiết kế chi tiết cơ sở dữ liệu PBL7")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    subtitle_run = subtitle.add_run(
        "Mỗi bảng dưới đây mô tả Field, Type và Note theo schema hiện tại của hệ thống."
    )
    subtitle_run.font.name = "Calibri"
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = RGBColor(85, 85, 85)

    summary = doc.add_paragraph()
    summary.add_run("Phạm vi: ").bold = True
    summary.add_run(
        f"{len(TABLES)} bảng chính trong schema: người dùng, giỏ hàng, sản phẩm, danh mục, đơn hàng, thanh toán và semantic search."
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    for table_def in TABLES:
        add_table_design(doc, table_def)

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT.resolve())
